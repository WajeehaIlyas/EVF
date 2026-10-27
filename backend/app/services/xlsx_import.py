"""Read a labelled-events spreadsheet (one sheet per feature) into plain records.

The expected layout is the team's "events" workbook:

    rows 1-5   a summary block (formulas, totals ...) - ignored
    a header row containing at least "Human Label", usually
               Track_ID | OCR / Plate | Date | Time | Day / Night | Human Label | Issue |
               OCR Correctness | Image Clarity | ... | MMC Status
    below      one row per event

Only the header text matters, not where it sits: the header row is found by looking for the
"Human Label" cell in the first rows, and columns are matched by name (case, spaces and
punctuation ignored), so extra or reordered columns are fine. Sheets named after a feature
("seatbelt", "mobilephone") are imported into that feature.

Nothing here touches the database; see app.api.v1.analytics for that.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

HEADER_SEARCH_ROWS = 30
MAX_SAMPLE_PROBLEMS = 5

_LABELS = {"tp": "TP", "truepositive": "TP", "fp": "FP", "falsepositive": "FP", "unsure": "Unsure"}
_TICK = {"correct", "yes", "y", "true", "ok", "right", "tick", "1"}
_CROSS = {"incorrect", "wrong", "no", "n", "false", "cross", "0"}


@dataclass
class ParsedRow:
    key: str                       # stable id of the row, used so re-importing never double counts
    label: str | None              # TP | FP | Unsure, or None while the row is still unlabelled
    vrm: str | None = None
    timestamp: datetime | None = None
    day_night: str | None = None
    track_id: str | None = None
    issue: str | None = None
    vrm_correct: bool | None = None
    make_correct: bool | None = None
    model_correct: bool | None = None
    color_correct: bool | None = None
    clear: bool | None = None      # image clarity: True clear, False unclear, None not stated
    make: str | None = None
    make_conf: float | None = None
    model: str | None = None
    model_conf: float | None = None
    color: str | None = None
    color_conf: float | None = None

    def columns(self) -> dict[str, Any]:
        """The values as ``archived_events`` columns (used to insert, update and compare)."""
        return {
            "human_label": self.label, "track_id": self.track_id, "vrm": self.vrm,
            "offence_timestamp": self.timestamp, "day_night": self.day_night, "issue": self.issue,
            "vrm_correct": self.vrm_correct, "make_correct": self.make_correct,
            "model_correct": self.model_correct, "color_correct": self.color_correct,
            "image_clarity_clear": self.clear,
            "vehicle_make": self.make, "make_confidence": self.make_conf,
            "vehicle_model": self.model, "model_confidence": self.model_conf,
            "vehicle_color": self.color, "color_confidence": self.color_conf,
        }


@dataclass
class SheetResult:
    name: str
    rows: list[ParsedRow] = field(default_factory=list)
    rows_read: int = 0
    unlabelled: int = 0            # rows with no Human Label yet: counted as violations, not as reviewed
    duplicates: int = 0            # the same event listed twice in the sheet
    invalid: int = 0               # a Human Label that is not TP / FP / Unsure
    problems: list[str] = field(default_factory=list)
    error: str | None = None       # the sheet could not be read at all (e.g. no header row)


def _norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower()) if value is not None else ""


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def parse_label(value: object) -> str | None:
    return _LABELS.get(_norm(value))


def parse_clarity(value: object) -> bool | None:
    """'Clear' -> True, 'Unclear' -> False, anything else -> not stated."""
    text = _norm(value)
    if text == "clear":
        return True
    if text in ("unclear", "notclear"):
        return False
    return None


def parse_tick(value: object) -> bool | None:
    """OCR Correctness column: 'Correct' / 'Incorrect' (also yes/no, true/false)."""
    text = _norm(value)
    if text in _TICK:
        return True
    if text in _CROSS:
        return False
    return None


def parse_mmc_status(value: object) -> tuple[bool | None, bool | None, bool | None]:
    """'MMC Status' text -> (make, model, colour) ticks.

    "all correct", "all wrong", or any mix such as "make correct", "make & colour correct" or
    "model correct, colour correct". When a status is written, anything it does not call correct
    counts as wrong. An empty status means the event was not MMC-checked: (None, None, None).
    """
    text = _text(value).lower()
    if not text:
        return None, None, None
    if "all correct" in text:
        return True, True, True
    if "all wrong" in text or "all incorrect" in text or "none correct" in text:
        return False, False, False

    def verdict(word: str) -> bool:
        if re.search(rf"{word}\s*(is\s*)?(wrong|incorrect)", text):
            return False
        return word in text and "correct" in text and f"{word} incorrect" not in text

    return verdict("make"), verdict("model"), verdict("colour") or verdict("color")


def _to_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10] if fmt.startswith("%Y") else text, fmt).date()
        except ValueError:
            continue
    return None


def _to_time(value: object) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    text = _text(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def _track(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _vehicle_text(value: object) -> str | None:
    """'Ford' -> 'Ford'; blank / 'N/A' / 'unknown' -> not stated."""
    text = _text(value)
    return None if _norm(text) in ("", "na", "none", "unknown", "null") else text[:100]


def _score(value: object) -> float | None:
    if isinstance(value, bool) or value is None or _text(value) == "":
        return None
    try:
        number = float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None   # the sheet writes 0.0 next to "N/A"


def _day_night(value: object) -> str | None:
    text = _norm(value)
    return "Day" if text == "day" else "Night" if text == "night" else None


def parse_sheet(name: str, rows: list[tuple[Any, ...]]) -> SheetResult:
    """Turn the rows of one worksheet into :class:`ParsedRow` records."""
    result = SheetResult(name=name)

    header_at = columns = None
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = [_norm(cell) for cell in row]
        if "humanlabel" in names:
            header_at, columns = index, {n: i for i, n in reversed(list(enumerate(names))) if n}
            break
    if header_at is None or columns is None:
        result.error = 'No header row with a "Human Label" column was found'
        return result

    def cell(row: tuple[Any, ...], *names: str) -> Any:
        for wanted in names:
            position = columns.get(wanted)
            if position is not None and position < len(row):
                return row[position]
        return None

    seen: set[str] = set()
    for offset, row in enumerate(rows[header_at + 1:], start=header_at + 2):
        if not any(_text(c) for c in row):
            continue
        result.rows_read += 1

        raw_label = cell(row, "humanlabel")
        label = parse_label(raw_label)
        if _text(raw_label) and label is None:
            result.invalid += 1
            if len(result.problems) < MAX_SAMPLE_PROBLEMS:
                result.problems.append(f'Row {offset}: "{_text(raw_label)}" is not TP, FP or Unsure')
            continue
        if label is None:
            result.unlabelled += 1  # still counted as a violation; a later import can label it

        day, moment = _to_date(cell(row, "date")), _to_time(cell(row, "time"))
        stamp = datetime.combine(day, moment or time(0, 0)) if day else None
        plate = _text(cell(row, "ocrplate", "plate", "vrm", "ocr")).upper().replace(" ", "") or None
        track = _track(cell(row, "trackid", "track"))
        when = stamp.strftime("%Y%m%d%H%M%S") if stamp else ""
        key = f"xlsx:{track}:{when}:{plate or ''}" if (track or when or plate) else f"xlsx:row{offset}"
        key = key[:100]
        if key in seen:
            result.duplicates += 1
            continue
        seen.add(key)

        make, model, colour = parse_mmc_status(cell(row, "mmcstatus")) if label else (None, None, None)
        result.rows.append(ParsedRow(
            key=key, label=label, vrm=plate, timestamp=stamp,
            day_night=_day_night(cell(row, "daynight")),
            track_id=track or None,
            issue=_text(cell(row, "issue"))[:500] or None,
            vrm_correct=parse_tick(cell(row, "ocrcorrectness")),
            make_correct=make, model_correct=model, color_correct=colour,
            clear=parse_clarity(cell(row, "imageclarity")) if label else None,
            make=_vehicle_text(cell(row, "vehiclebrand", "vehiclemake")),
            make_conf=_score(cell(row, "vehiclebrandscore", "vehiclemakescore")),
            model=_vehicle_text(cell(row, "vehiclemodel")),
            model_conf=_score(cell(row, "vehiclemodelscore")),
            color=_vehicle_text(cell(row, "vehiclecolor", "vehiclecolour")),
            color_conf=_score(cell(row, "vehiclecolorscore", "vehiclecolourscore")),
        ))
    return result


def read_workbook(data: bytes) -> list[SheetResult]:
    """Parse every worksheet of an .xlsx file (given as bytes)."""
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        return [parse_sheet(sheet.title, list(sheet.iter_rows(values_only=True)))
                for sheet in workbook.worksheets]
    finally:
        workbook.close()
